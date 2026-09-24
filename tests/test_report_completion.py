"""Report snapshots, review guards, audience privacy and editable exports."""
from __future__ import annotations

import copy
import csv
import io
import json
import os
import zipfile
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

import openpyxl
import pytest

from apps.api.snapshot_store import unpack as unpack_snapshot
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pypdf import PdfReader
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from apps.api import identity
from apps.api.db import Base, Document, FindingRow, Project, ReportRow, VerificationRun, get_db
from apps.api.routers import reports
from apps.api.schemas import ReportRequest
from apps.api.workspace import CaseIssue, ClaimAssessment, FindingWorkflow, ReportReview, ReviewRevision
from packages.common.storage import LocalObjectStorage, sha256_bytes
from packages.common.terminology import TERMINOLOGY_VERSION, term_label
from packages.report_engine.snapshot import canonical_hash


@pytest.fixture()
def report_case(tmp_path, monkeypatch):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    session = Session(engine, expire_on_commit=False)
    storage = LocalObjectStorage(tmp_path / "storage")
    monkeypatch.setattr(reports, "get_storage", lambda: storage)
    audit = SimpleNamespace(manifest=lambda **kw: {"event_count": 0, "chain_valid": True,
                                                  "head_hash": "audit-head", "generated_at": "2026-01-01"}, record=Mock())
    monkeypatch.setattr(reports, "make_audit", lambda _: audit)
    project = Project(id="p1", name="Report case", scope_revision=1)
    other = Project(id="p2", name="Another project", scope_revision=0)
    session.add_all([project, other])
    session.flush()
    source = b"source bytes"
    document = Document(id="d1", project_id="p1", filename="case.pdf", sha256=sha256_bytes(source),
                        storage_key=storage.put_original("case", source), included_in_verification=True)
    session.add(document)
    session.flush()
    data = {"finding_id": "f1", "type": "CASE_NOT_FOUND", "status": "UNVERIFIED", "severity": "HIGH",
            "evidence_grade": "U", "confidence": 0.2, "document_id": "d1", "page": 1, "block_id": "b1",
            "title": "=HYPERLINK(\"https://invalid.example\")", "detail": "Unverified finding <script>literal</script>",
            "advisory_only": False, "engine": "legal", "has_sealed_content": True,
            "sealed_excerpt": "DO_NOT_EXPORT_SEALED", "review_note": "", "review_status": "NEEDS_REVIEW",
            "evidence": [{"sealed": True, "excerpt": "DO_NOT_EXPORT_SEALED"}], "sources": []}
    input_snapshot = {"scope_revision": 1, "project": {"name": "Report case", "parties": ["홍길동"]},
                      "context": {"profile": "STANDARD", "case_date": None, "external_ai_policy": "LOCAL_ONLY"},
                      "documents": [{"document_id": "d1", "sha256": document.sha256}]}
    document_result = {"document_id": "d1", "filename": "case.pdf", "sha256": document.sha256,
                       "parser": "fixture", "quarantined": True, "rag_indexable": False,
                       "warnings": [], "authorship": {}, "masked_preview": {}, "entities": [], "events": [],
                       "citations": [], "claims": [{"claim_id": "c1", "text": "홍길동 user@example.com 주장"}],
                       "engine_data": {"page_coverage": [{"page": 1, "status": "unsupported"}],
                                       "unavailable_stages": ["signature verification unavailable"],
                                       "legal_verdicts": [{"citation_id": "citation1", "levels": {"temporal": "UNVERIFIED"}}]},
                       "source_records": [{"adapter": "fixture", "url": "https://official.example/full/source",
                                           "response_hash": "source-digest", "retrieved_at": "2026-01-01",
                                           "status": "TIMEOUT", "note": "full source note"}],
                       "pages": [{"page_number": 1, "blocks": [{"text": "raw-private-page"}]}], "findings": [data]}
    run = VerificationRun(id="r1", project_id="p1", document_ids=["d1"], state="PARTIAL_COMPLETED",
                          started_at=datetime(2026, 1, 1), finished_at=datetime(2026, 1, 1, 1),
                          verification_key="key", input_snapshot=input_snapshot,
                          unavailable_sources=[{"name": "law", "status": "MISSING_KEY", "note": "not executed"}],
                          unverified_items=[{"kind": "citation", "raw_text": "citation", "reason": "source unavailable"}],
                          result_json={"documents": [document_result], "project_findings": [], "model_executions": []})
    session.add(run)
    session.flush()
    finding = FindingRow(id="f1", project_id="p1", run_id="r1", document_id="d1", type="CASE_NOT_FOUND",
                         status="UNVERIFIED", severity="HIGH", evidence_grade="U", confidence=0.2,
                         title=data["title"], detail=data["detail"], data=data, sealed_excerpt="DO_NOT_EXPORT_SEALED")
    issue = CaseIssue(id="i1", project_id="p1", title="계약상 책임", elements=["계약 성립"], updated_by="reviewer")
    session.add_all([finding, issue])
    session.commit()
    app = FastAPI()
    app.include_router(reports.router, prefix="/api")

    def database():
        yield session
    app.dependency_overrides[get_db] = database

    @app.middleware("http")
    async def authenticated(request, call_next):
        principal = identity.Principal("outsider", "unrelated-org", "MEMBER", "token") if request.headers.get("x-test-outsider") else identity.Principal("local-owner", None, "ADMIN", "local")
        token = identity._principal.set(principal)
        try:
            return await call_next(request)
        finally:
            identity._principal.reset(token)

    with TestClient(app) as client:
        yield SimpleNamespace(client=client, session=session, project=project, run=run,
                              document=document, finding=finding, storage=storage, tmp_path=tmp_path)
    session.close()
    engine.dispose()


def create(case, formats=None, audience="INTERNAL"):
    payload = {"audience": audience}
    if formats is not None:
        payload["formats"] = formats
    response = case.client.post("/api/projects/p1/reports", json=payload)
    assert response.status_code == 201, response.text
    report = response.json()
    assert not any("error" in a for a in report["artifacts"].values()), report
    return report


def finalize(case, draft, **values):
    return case.client.post(f"/api/reports/{draft['report_id']}/finalize", json={
        "note": "Reviewed with recorded limitations", "acknowledge_unresolved": True,
        "acknowledge_stale_scope": True, "acknowledge_privacy": True, **values})


def download(case, report, fmt):
    response = case.client.get(f"/api/reports/{report['report_id']}/download/{fmt}")
    assert response.status_code == 200, response.text
    return response.content


def test_default_formats_and_schema_validation():
    assert ReportRequest().formats == ["pdf", "xlsx", "csv", "json", "manifest"]
    assert ReportRequest(audience="shareable").audience == "SHAREABLE"
    for formats in ([], ["docx", "docx"], ["unknown"]):
        with pytest.raises(ValueError):
            ReportRequest(formats=formats)
    assert TERMINOLOGY_VERSION == "2"
    assert term_label("position", "UNKNOWN") == "부지"
    assert term_label("position", "ALTERNATIVE") == "예비적 주장"


def test_terminology_catalog_requires_project_access(report_case):
    response = report_case.client.get("/api/projects/p1/report-terminology")
    assert response.status_code == 200 and response.json()["version"] == "2"
    assert response.json()["groups"]["verification_status"]["UNVERIFIED"]["label"] == "미검증"
    assert report_case.client.get("/api/projects/p1/report-terminology", headers={"x-test-outsider": "yes"}).status_code == 404


def test_finalize_requires_acknowledgments_and_preserves_draft(report_case):
    case = report_case
    draft = create(case, ["json"])
    original_bytes = download(case, draft, "json")
    original_review = copy.deepcopy(case.session.get(ReportReview, draft["report_id"]).review_snapshot)
    case.project.scope_revision += 1
    case.session.commit()
    preflight = case.client.get(f"/api/reports/{draft['report_id']}/preflight").json()
    assert preflight["incomplete_reviews"] and preflight["stale_scope"] and preflight["sensitive_concerns"]
    assert any(item["type"] == "PERSONAL_DATA" for item in preflight["sensitive_concerns"])
    for name in ("acknowledge_unresolved", "acknowledge_stale_scope", "acknowledge_privacy"):
        rejected = finalize(case, draft, **{name: False})
        assert rejected.status_code == 409
        assert name in rejected.json()["detail"]["missing_acknowledgments"]
    invalid = finalize(case, draft, acknowledge_privacy="true")
    assert invalid.status_code == 422
    case.session.add(FindingWorkflow(finding_id="f1", workflow_state="COMPLETED", decision="AGREED",
                                    note="human snapshot note", updated_by="reviewer"))
    case.session.add(ClaimAssessment(project_id="p1", run_id="r1", document_id="d1", claim_id="c1", issue_id="i1",
                                    position="DENIED", support_status="INSUFFICIENT", updated_by="reviewer",
                                    evidence_links=[{"document_id": "d1", "page": 1, "relation": "CONTEXT"}]))
    case.session.add(ReviewRevision(project_id="p1", subject_type="FINDING", subject_id="f1", actor="reviewer",
                                   after={"workflow_state": "COMPLETED"}))
    case.session.commit()
    stale_confirmation = finalize(case, draft, expected_preflight_hash=preflight["preflight_hash"])
    assert stale_confirmation.status_code == 409
    final_response = finalize(case, draft)
    assert final_response.status_code == 201, final_response.text
    final = final_response.json()
    assert final["report_id"] != draft["report_id"] and final["state"] == "FINAL"
    assert final["created_by"] == final["finalized_by"] == "local-owner" and final["finalized_at"]
    frozen = json.loads(download(case, final, "json"))
    assert frozen["documents"][0]["findings"][0]["status"] == "UNVERIFIED"
    assert frozen["review_snapshot"]["workflow"][0]["workflow_state"] == "COMPLETED"
    assert frozen["review_snapshot"]["matrix"]["claims"][0]["assessment"]["position"] == "DENIED"
    assert frozen["review_snapshot"]["review_history"][0]["actor"] == "reviewer"
    assert case.finding.status == "UNVERIFIED"
    final_bytes = download(case, final, "json")
    case.session.get(FindingWorkflow, "f1").note = "later edits"
    case.run.result_json = {"documents": []}
    case.session.commit()
    assert download(case, final, "json") == final_bytes
    assert download(case, draft, "json") == original_bytes
    assert case.session.get(ReportReview, draft["report_id"]).review_snapshot == original_review
    assert finalize(case, final).status_code == 409


def test_final_uses_frozen_engine_even_if_run_changes(report_case):
    case = report_case
    draft = create(case, ["json"])
    changed = copy.deepcopy(case.run.result_json)
    changed["documents"][0]["findings"][0]["status"] = "VERIFIED"
    changed["documents"][0]["filename"] = "changed-live-name.pdf"
    case.run.result_json = changed
    case.session.commit()
    final = finalize(case, draft).json()
    payload = json.loads(download(case, final, "json"))
    assert payload["documents"][0]["findings"][0]["status"] == "UNVERIFIED"
    assert payload["documents"][0]["filename"] == "case.pdf"


def test_cross_project_and_unauthorized_report_access(report_case):
    case = report_case
    assert case.client.post("/api/projects/p2/reports", json={"run_id": "r1", "formats": ["json"]}).status_code == 404
    draft = create(case, ["json"])
    for url in (f"/api/reports/{draft['report_id']}", f"/api/reports/{draft['report_id']}/preflight",
                f"/api/reports/{draft['report_id']}/download/json", "/api/projects/p1/reports"):
        assert case.client.get(url, headers={"x-test-outsider": "yes"}).status_code == 404
    assert case.client.post(f"/api/reports/{draft['report_id']}/finalize", headers={"x-test-outsider": "yes"},
                            json={"note": "outsider", "acknowledge_privacy": True}).status_code == 404


def test_shareable_removes_notes_pages_sealed_and_pii(report_case):
    case = report_case
    case.session.add(FindingWorkflow(finding_id="f1", note="PRIVATE_INTERNAL_NOTE", updated_by="reviewer"))
    case.session.commit()
    draft = create(case, ["json", "docx"], "shareable")
    output = download(case, draft, "json").decode()
    xml = zipfile.ZipFile(io.BytesIO(download(case, draft, "docx"))).read("word/document.xml").decode()
    for value in ("DO_NOT_EXPORT_SEALED", "PRIVATE_INTERNAL_NOTE", "raw-private-page", "user@example.com", "홍길동"):
        assert value not in output and value not in xml
    assert "SHAREABLE" in output
    assert "full source note" in output and "full source note" in xml
    private = case.session.get(ReportReview, draft["report_id"]).review_snapshot
    assert private["workflow"][0]["note"] == "PRIVATE_INTERNAL_NOTE"
    assert case.client.post("/api/projects/p1/reports", json={"audience": "SHAREABLE", "formats": ["highlight"]}).status_code == 422


def test_cross_format_labels_appendix_and_docx_zip(report_case):
    case = report_case
    draft = create(case, ["pdf", "xlsx", "csv", "json", "manifest", "docx"])
    response = finalize(case, draft)
    assert response.status_code == 201, response.text
    final = response.json()
    exports = {fmt: download(case, final, fmt) for fmt in final["formats"]}
    payload = json.loads(exports["json"])
    label = payload["report"]["label"]
    assert label == final["label"] and payload["report"]["state"] == "FINAL"
    pdf_text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(exports["pdf"])).pages)
    workbook = openpyxl.load_workbook(io.BytesIO(exports["xlsx"]))
    xlsx_text = "\n".join(str(c.value) for ws in workbook for row in ws for c in row if c.value is not None)
    archive = zipfile.ZipFile(io.BytesIO(exports["docx"]))
    word = archive.read("word/document.xml").decode()
    for text in (pdf_text, xlsx_text, word):
        assert "FINAL" in text and final["snapshot_hash"] in text and "UNVERIFIED" in text
        assert "signature verification unavailable" in text and "full source note" in text
        assert "official.example/full/source" in text
        assert "DO_NOT_EXPORT_SEALED" not in text
    assert label in xlsx_text and label in word
    assert not any("vbaProject" in name or "embeddings/" in name for name in archive.namelist())
    assert "<w:instrText" not in word and "<w:altChunk" not in word
    assert "&lt;script&gt;literal&lt;/script&gt;" in word
    assert not any('TargetMode="External"' in archive.read(name).decode() for name in archive.namelist() if name.endswith(".rels"))
    assert payload["report"]["editable_copy_notice"] in word
    assert json.loads(exports["manifest"])["report"]["label"] == label
    csv_rows = list(csv.DictReader(io.StringIO(exports["csv"].decode("utf-8-sig"))))
    assert csv_rows[0]["report_label"] == label
    qa_dir = os.getenv("ACAS_REPORT_QA_DIR") or str(case.tmp_path)
    if qa_dir:
        output = Path(qa_dir)
        output.mkdir(parents=True, exist_ok=True)
        (output / "report.docx").write_bytes(exports["docx"])
        (output / "report.pdf").write_bytes(exports["pdf"])


def test_insufficient_claim_support_still_requires_unresolved_acknowledgment(report_case):
    case = report_case
    changed = copy.deepcopy(case.run.result_json)
    changed["documents"][0]["findings"][0]["status"] = "VERIFIED"
    case.run.result_json = changed
    case.run.unverified_items = []
    case.run.unavailable_sources = []
    case.run.state = "COMPLETED"
    case.session.add(FindingWorkflow(finding_id="f1", workflow_state="COMPLETED", decision="AGREED", updated_by="reviewer"))
    case.session.add(ClaimAssessment(project_id="p1", run_id="r1", document_id="d1", claim_id="c1", issue_id="i1",
                                    position="DENIED", support_status="INSUFFICIENT", updated_by="reviewer"))
    case.session.commit()
    draft = create(case, ["json"])
    checks = case.client.get(f"/api/reports/{draft['report_id']}/preflight").json()
    assert checks["unresolved_claims"] and not checks["unresolved_findings"]
    assert finalize(case, draft, acknowledge_unresolved=False).status_code == 409


def test_formula_injection_and_long_results_are_preserved(report_case):
    case = report_case
    changed = copy.deepcopy(case.run.result_json)
    finding = changed["documents"][0]["findings"][0]
    finding["title"] = " \t=HYPERLINK(\"https://invalid.example\")"
    finding["detail"] = "@SUM(1+1)" + "x" * 40000 + "FULL_TEXT_END"
    case.run.result_json = changed
    case.run.unverified_items = [{"kind": "citation", "raw_text": f"+unverified-{n}", "reason": f"reason-{n}"} for n in range(65)]
    case.session.commit()
    draft = create(case, ["xlsx", "csv", "pdf"])
    workbook = openpyxl.load_workbook(io.BytesIO(download(case, draft, "xlsx")))
    assert "Long values" in workbook.sheetnames
    assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)
    all_text = "".join(str(c.value or "") for ws in workbook for row in ws for c in row)
    assert "FULL_TEXT_END" in all_text and "reason-64" in all_text
    csv_rows = list(csv.DictReader(io.StringIO(download(case, draft, "csv").decode("utf-8-sig"))))
    assert csv_rows[0]["title"].startswith("'") and csv_rows[0]["detail"].startswith("'")
    assert csv_rows[0]["detail"].endswith("FULL_TEXT_END")
    pdf_text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(download(case, draft, "pdf"))).pages)
    # 표 칸의 긴 값은 줄바꿈되어 추출되므로 줄바꿈을 지우고 끝까지 실렸는지 본다.
    assert "reason-64" in pdf_text and "FULL_TEXT_END" in pdf_text.replace("\n", "")


def test_unavailable_stages_require_ack_even_without_unverified_findings(report_case):
    case = report_case
    changed = copy.deepcopy(case.run.result_json)
    changed["documents"][0]["findings"][0]["status"] = "VERIFIED"
    changed["documents"][0]["claims"] = []
    case.run.result_json = changed
    case.run.state = "COMPLETED"
    case.run.unverified_items = []
    case.run.unavailable_sources = []
    case.session.add(FindingWorkflow(finding_id="f1", workflow_state="COMPLETED", decision="AGREED", updated_by="reviewer"))
    case.session.commit()
    draft = create(case, ["json"])
    checks = case.client.get(f"/api/reports/{draft['report_id']}/preflight").json()
    assert checks["unavailable_stages"] and not checks["unresolved_findings"] and not checks["unresolved_claims"]
    assert finalize(case, draft, acknowledge_unresolved=False).status_code == 409


def test_failed_final_render_does_not_commit_final(report_case, monkeypatch):
    case = report_case
    draft = create(case, ["json"])
    before = download(case, draft, "json")
    def failure(*args, **kwargs):
        raise RuntimeError("renderer unavailable")
    monkeypatch.setattr(reports, "_render", failure)
    assert finalize(case, draft).status_code == 409
    assert len(list(case.session.scalars(select(ReportRow)))) == 1
    assert download(case, draft, "json") == before


def test_snapshot_and_artifact_tampering_are_rejected(report_case):
    case = report_case
    draft = create(case, ["json"])
    review = case.session.get(ReportReview, draft["report_id"])
    assert canonical_hash(unpack_snapshot(case.storage, review.review_snapshot)) == review.snapshot_hash
    review.review_snapshot = {**review.review_snapshot, "unexpected": True}
    case.session.commit()
    assert finalize(case, draft).status_code == 409
    artifact = draft["artifacts"]["json"]
    case.storage.path(artifact["storage_key"]).write_bytes(b"tampered")
    assert case.client.get(artifact["download"]).status_code == 409


def test_reports_use_scoped_manifest_and_do_not_label_unchecked_chain(report_case, monkeypatch):
    calls = []

    def scoped_manifest(project_id, session):
        calls.append(project_id)
        return {"project_id": project_id, "event_count": 0, "events": [],
                "chain_valid": None, "event_hashes_valid": True,
                "verification_scope": "project_events", "head_hash": "scoped-head"}

    monkeypatch.setattr(reports, "project_manifest", scoped_manifest)
    draft = create(report_case, ["pdf", "manifest"])
    response = finalize(report_case, draft)
    assert response.status_code == 201, response.text
    final = response.json()
    assert calls == ["p1", "p1"]
    manifest = json.loads(download(report_case, final, "manifest"))
    assert manifest["chain_valid"] is None
    assert manifest["verification_scope"] == "project_events"
    pdf_text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(download(report_case, final, "pdf"))).pages)
    assert "미검증 (사건별 기록만 조회)" in pdf_text


def test_report_rendering_does_not_hold_a_database_writer_lock(report_case, monkeypatch):
    from sqlalchemy import event

    writes, observed = [], []
    engine = report_case.session.get_bind()
    original = reports._render

    def track(connection, cursor, statement, parameters, context, executemany):
        if statement.lstrip().split(" ", 1)[0].upper() in {"INSERT", "UPDATE", "DELETE"}:
            writes.append(statement)

    def render(*args, **kwargs):
        assert not writes, "Report rendering must precede database writes"
        observed.append(args[0])
        return original(*args, **kwargs)

    event.listen(engine, "before_cursor_execute", track)
    monkeypatch.setattr(reports, "_render", render)
    try:
        draft = create(report_case, ["json", "manifest"])
        assert writes
        writes.clear()
        response = finalize(report_case, draft)
        assert response.status_code == 201, response.text
        assert writes and observed == ["json", "manifest", "json", "manifest"]
    finally:
        event.remove(engine, "before_cursor_execute", track)


def _synchronous_jobs(case, monkeypatch):
    """생성 스레드 대신 요청 안에서 바로 실행해 진행 기록을 순서대로 모은다."""
    engine = case.session.get_bind()
    monkeypatch.setattr(reports, "_job_session", lambda: Session(engine, expire_on_commit=False))
    monkeypatch.setattr(reports, "_launch_job", reports._run_report_job)
    seen = []
    original = reports._job_update

    def record(job_id, **values):
        seen.append(dict(values))
        original(job_id, **values)
    monkeypatch.setattr(reports, "_job_update", record)
    return seen


def test_report_job_shows_progress_until_the_report_exists(report_case, monkeypatch):
    case = report_case
    seen = _synchronous_jobs(case, monkeypatch)
    started = case.client.post("/api/projects/p1/report-jobs", json={"formats": ["pdf", "docx", "json"]})
    assert started.status_code == 202, started.text
    job = started.json()
    assert job["state"] == "QUEUED" and job["percent"] == 0
    polled = case.client.get(f"/api/projects/p1/report-jobs/{job['job_id']}").json()
    assert polled["state"] == "COMPLETED" and polled["percent"] == 100 and polled["error"] is None
    reports_list = case.client.get("/api/projects/p1/reports").json()
    report = next(item for item in reports_list if item["report_id"] == polled["report_id"])
    assert set(report["artifacts"]) == {"pdf", "docx", "json"}
    assert not any("error" in artifact for artifact in report["artifacts"].values())
    # 단계마다 이름과 진행률을 남기고, 진행률은 되돌아가지 않는다.
    stages = [item["stage"] for item in seen if item.get("stage")]
    assert any("PDF 작성 중 (1/3)" == stage for stage in stages)
    assert any("Word 작성 중 (2/3)" == stage for stage in stages)
    percents = [item["percent"] for item in seen if "percent" in item]
    assert percents == sorted(percents) and percents[-1] == 100
    assert case.client.get("/api/projects/p1/report-jobs?active=true").json() == []
    assert case.client.get(f"/api/projects/p1/report-jobs/{job['job_id']}",
                           headers={"x-test-outsider": "yes"}).status_code == 404


def test_report_job_rejects_bad_requests_before_queueing(report_case, monkeypatch):
    case = report_case
    _synchronous_jobs(case, monkeypatch)
    assert case.client.post("/api/projects/p1/report-jobs",
                            json={"audience": "SHAREABLE", "formats": ["highlight"]}).status_code == 422
    assert case.client.post("/api/projects/p1/report-jobs", json={"formats": ["pptx"]}).status_code == 422
    assert case.client.get("/api/projects/p1/report-jobs").json() == []


def test_report_job_failure_is_explained(report_case, monkeypatch):
    case = report_case
    _synchronous_jobs(case, monkeypatch)

    def broken(*args, **kwargs):
        raise RuntimeError("renderer crashed")
    monkeypatch.setattr(reports, "_build_draft", broken)
    job = case.client.post("/api/projects/p1/report-jobs", json={"formats": ["pdf"]}).json()
    polled = case.client.get(f"/api/projects/p1/report-jobs/{job['job_id']}").json()
    assert polled["state"] == "FAILED" and polled["report_id"] is None
    assert "보고서를 만들지 못했습니다 (RuntimeError)" in polled["error"]


def test_interrupted_report_job_is_not_left_running(report_case):
    from datetime import timedelta
    from apps.api.workspace import ReportJob
    case = report_case
    case.session.add(ReportJob(id="rjb_stale", project_id="p1", run_id="r1", created_by="local-owner",
                               request={"formats": ["pdf"]}, state="RUNNING", stage="PDF 작성 중 (1/1)", percent=40,
                               created_at=datetime.utcnow() - timedelta(minutes=10),
                               updated_at=datetime.utcnow() - timedelta(minutes=5)))
    case.session.commit()
    polled = case.client.get("/api/projects/p1/report-jobs/rjb_stale").json()
    assert polled["state"] == "FAILED" and "다시 시작" in polled["error"]


def test_technical_appendix_lists_each_claim_text_once():
    from packages.report_engine.snapshot import CLAIM_REFERENCE_NOTE, compact_claim_rows
    claim = {"claim_id": "c1", "text": "긴 주장 전문"}
    result = {"documents": [{"claims": [claim]}],
              "review_snapshot": {"matrix": {"claims": [{"claim": dict(claim), "review_status": "PARTIAL"}]},
                                  "preflight": {"incomplete_claim_reviews": [{"claim": dict(claim)}],
                                                "unresolved_claims": [{"claim": dict(claim)}]}}}
    compact_claim_rows(result)
    assert json.dumps(result, ensure_ascii=False).count("긴 주장 전문") == 1
    assert result["review_snapshot"]["matrix"]["claims"][0] == {
        "claim": {"claim_id": "c1", "note": CLAIM_REFERENCE_NOTE}, "review_status": "PARTIAL"}


def test_word_and_pdf_appendix_scale_with_many_entries(report_case, monkeypatch):
    """부록 항목이 많아도 Word 문단은 구역 설정 앞에 순서대로, PDF 줄은 페이지 폭 안에 들어간다."""
    import time
    from docx import Document as WordDocument
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from packages.report_engine import docx_report
    from packages.report_engine.pdf_report import _appendix_blocks, _styles
    style = _styles()["small"]
    entries = [(f"documents[0].claims[{i}].text", "가나다 " * 40 + "https://official.example/full/source")
               for i in range(3000)]
    started = time.perf_counter()
    blocks = _appendix_blocks(entries, style, 470)
    assert time.perf_counter() - started < 10
    lines = [line for block in blocks for line in block.lines]
    assert max(stringWidth(line, style.fontName, style.fontSize) for line in lines) <= 470
    assert "".join(line.strip() for line in lines).count("https://official.example/full/source") == 3000

    monkeypatch.setattr(docx_report, "technical_payload", lambda _: {"items": [f"항목 {i}" for i in range(5000)]})
    started = time.perf_counter()
    response = report_case.client.post("/api/projects/p1/reports", json={"formats": ["docx"], "detail_level": "FULL"})
    assert response.status_code == 201, response.text
    draft = response.json()
    assert time.perf_counter() - started < 20
    word = WordDocument(io.BytesIO(download(report_case, draft, "docx")))
    texts = [p.text for p in word.paragraphs]
    assert texts.index("항목 0") < texts.index("항목 4999")
    assert word.element.body[-1].tag.endswith("sectPr")


def test_summary_report_keeps_review_content_and_moves_the_full_record_to_json(report_case):
    """요약본은 판단 근거(수행하지 못한 단계·실패한 조회·미확인 항목)를 싣고 원자료 나열은 JSON에 둔다."""
    case = report_case
    summary = case.client.post("/api/projects/p1/reports", json={"formats": ["pdf", "docx", "xlsx"]}).json()
    assert summary["formats"] == ["pdf", "docx", "xlsx", "json"]  # 전체 기록을 담을 JSON을 함께 만든다
    full = case.client.post("/api/projects/p1/reports",
                            json={"formats": ["pdf", "docx", "xlsx"], "detail_level": "FULL"}).json()
    assert "json" not in full["formats"]
    texts = {}
    for name, report in (("summary", summary), ("full", full)):
        pdf = download(case, report, "pdf")
        texts[name] = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
        word = zipfile.ZipFile(io.BytesIO(download(case, report, "docx"))).read("word/document.xml").decode()
        workbook = openpyxl.load_workbook(io.BytesIO(download(case, report, "xlsx")))
        sheet_text = "\n".join(str(c.value) for ws in workbook for row in ws for c in row if c.value is not None)
        for text in (texts[name], word, sheet_text):
            assert "signature verification unavailable" in text and "full source note" in text
            assert "official.example/full/source" in text and "UNVERIFIED" in text
        texts[name + "_sheets"] = workbook.sheetnames
    assert "전체 기술 기록(실행 당시의 모든 필드)을 싣지 않았습니다" in texts["summary"]
    assert "검증근거" in texts["summary_sheets"] and "기술부록" not in texts["summary_sheets"]
    assert "기술부록" in texts["full_sheets"]
    assert "input_snapshot.context.profile" in texts["full"] and "input_snapshot.context.profile" not in texts["summary"]
    payload = json.loads(download(case, summary, "json"))
    assert payload["report"]["detail_level"] == "SUMMARY" and payload["documents"][0]["claims"]


def test_finalized_report_keeps_the_draft_detail_level(report_case):
    case = report_case
    draft = case.client.post("/api/projects/p1/reports", json={"formats": ["pdf"], "detail_level": "FULL"}).json()
    final = finalize(case, draft).json()
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(download(case, final, "pdf"))).pages)
    assert "detail_level: FULL" in text and "15. 전체 기술 기록" in text


def test_summary_verdict_table_has_a_document_column_grouped_by_document(report_case):
    """v3 D8: 요약 판정 표의 모든 항목에 문서명을 적고 문서별로 묶는다."""
    from docx import Document as WordDocument
    report = create(report_case, ["docx"])
    word = WordDocument(io.BytesIO(download(report_case, report, "docx")))
    table = word.tables[0]
    header = [c.text for c in table.rows[0].cells]
    assert header[0] == "문서"
    rows = [[c.text for c in r.cells] for r in table.rows[1:]]
    assert rows and all(r[0] for r in rows)
    assert report_case.document.filename in {r[0] for r in rows}


def test_citation_error_table_merges_the_verdict_into_the_basis_column(report_case):
    """주장 평가는 인용 오류·미확인 근거 칸에 함께 싣고, 법리 검토 칸을 가장 넓게 둔다(Word·PDF)."""
    from docx import Document as WordDocument
    case = report_case
    result = dict(case.run.result_json)
    result["documents"] = [{**result["documents"][0], "ai_hallucination_table": [{
        "location": "3쪽", "claim_text": "처분은 위법하다", "cited_authority": "대법원 2099. 1. 1. 선고 2099두1 판결",
        "ai_generation_basis": "공식 DB에서 확인되지 않음", "legal_reasoning": "법리 검토 본문",
        "recommended_counteraction": "원문 제출 요구", "validity_verdict": "근거 결여"}]}]
    case.run.result_json = result
    case.session.commit()
    report = create(case, ["docx", "pdf"])
    word = WordDocument(io.BytesIO(download(case, report, "docx")))
    [table] = [t for t in word.tables if t.rows[0].cells[-1].text == "법리적 타당성 검토 및 반박 근거"]
    header = [c.text for c in table.rows[0].cells]
    assert header == ["위치", "문서 주장 / 인용", "인용 오류·미확인 근거 및 주장 평가", "법리적 타당성 검토 및 반박 근거"]
    basis = table.rows[1].cells[2].text
    assert "[평가] 근거 결여" in basis and "공식 DB에서 확인되지 않음" in basis
    widths = [c.width for c in table.rows[0].cells]
    assert widths[3] == max(widths)
    # 칸 안 줄바꿈 위치와 가운뎃점 글리프(·/・)는 설치된 글꼴에 따라 달라지므로 공백·가운뎃점을 빼고 비교한다.
    text = "".join("".join((p.extract_text() or "").split())
                   for p in PdfReader(io.BytesIO(download(case, report, "pdf"))).pages)
    text = text.replace("·", "").replace("・", "")
    assert "평가]근거결여" in text and "인용오류미확인근거및주장평가" in text


# --- 고정본의 큰 부분은 DB 밖(파일 저장소)에 내용 주소로 보관한다 --------------------------
def test_large_snapshot_parts_live_outside_the_database_and_are_shared(report_case):
    """검증 결과·사전 점검은 파일로 옮기고 DB에는 참조만 둔다. 같은 내용은 한 번만 저장한다."""
    import json as _json

    from apps.api.snapshot_store import BLOB_KEYS, MARKER
    case = report_case
    first, second = create(case, ["json"]), create(case, ["json"])
    stored = [case.session.get(ReportReview, r["report_id"]).review_snapshot for r in (first, second)]
    for snapshot in stored:
        for key in BLOB_KEYS:
            assert MARKER in snapshot[key]
        assert len(_json.dumps(snapshot, default=str)) < 20000      # DB에는 참조와 작은 항목만
    # 같은 검증 결과는 같은 파일을 가리킨다.
    assert stored[0]["engine_result"] == stored[1]["engine_result"]
    assert finalize(case, first).status_code == 201
    assert case.client.get(first["artifacts"]["json"]["download"]).status_code == 200


def test_tampered_or_missing_snapshot_file_is_rejected(report_case):
    case = report_case
    draft = create(case, ["json"])
    ref = case.session.get(ReportReview, draft["report_id"]).review_snapshot["engine_result"]
    import gzip
    case.storage.path(ref["$snapshot_blob"]).write_bytes(gzip.compress(b'{"documents": []}'))
    response = finalize(case, draft)
    assert response.status_code == 409 and "고정본" in str(response.json())
    case.storage.path(ref["$snapshot_blob"]).unlink()
    assert finalize(case, draft).status_code == 409


def test_legacy_inline_snapshot_is_still_readable(report_case):
    """예전 형식(전부 DB에 있는 고정본)으로 저장된 보고서도 그대로 확정·내려받기가 된다."""
    case = report_case
    draft = create(case, ["json"])
    review = case.session.get(ReportReview, draft["report_id"])
    review.review_snapshot = unpack_snapshot(case.storage, review.review_snapshot)
    case.session.commit()
    assert canonical_hash(review.review_snapshot) == review.snapshot_hash
    assert finalize(case, draft).status_code == 201

