"""'법리적 타당성 검토 및 반박 근거' 칸의 항목별 줄바꿈(검토 결과 / AI 교차검증 요약 / 모델별 의견)."""
from __future__ import annotations

import io
import os
from pathlib import Path
from types import SimpleNamespace
from urllib.parse import urlsplit

import pytest

from packages.legal_engine.argument_validity_verifier import HallucinationTableRow
from packages.legal_engine.reasoning_format import format_reasoning, reasoning_sections, row_cell_text

ROOT = Path(__file__).resolve().parents[1]
OPINIONS = [{"provider": "anthropic", "verdict": "부당", "reasoning": "요건 사실이 없다"},
            {"provider": "openai", "verdict": "타당", "reasoning": ""},
            {"provider": "gemini", "verdict": "부당", "reasoning": "판시 취지와 다르다"}]


@pytest.mark.parametrize("review, label, opinions, expected", [
    ("근거가 확장되지 않습니다.", "3개 모델 의견 불일치 — 직접 검토 필요", OPINIONS,
     ["타당성 검토 결과 : 근거가 확장되지 않습니다.", "AI 교차검증 결과 참고 : 3개 모델 의견 불일치 — 직접 검토 필요",
      "  (Anthropic : 부당)", "      근거: 요건 사실이 없다", "  (OpenAI : 타당)", "  (Gemini : 부당)",
      "      근거: 판시 취지와 다르다"]),
    ("사건번호 형식이 성립하지 않습니다.", None, [], ["타당성 검토 결과 : 사건번호 형식이 성립하지 않습니다."]),
    ("공식 DB에서 확인하지 못했습니다.", "1개 모델(gemini) 의견(교차검증 아님)",
     [{"provider": "gemini", "verdict": "판단 불가"}],
     ["타당성 검토 결과 : 공식 DB에서 확인하지 못했습니다.", "AI 교차검증 결과 참고 : 1개 모델(gemini) 의견(교차검증 아님)",
      "  (Gemini : 판단 불가)"]),
])
def test_each_part_is_on_its_own_line(review, label, opinions, expected):
    assert format_reasoning(review, label, opinions).split("\n") == expected


def test_row_dict_carries_structure_and_formatted_text():
    row = HallucinationTableRow(location="2쪽", claim_text="c", cited_authority="대법원 2020다1 판결", authority_exists=False,
                                ai_generation_basis="b", validity_verdict="부당", legal_reasoning="법리상 부당",
                                recommended_counteraction="원문 확인")
    plain = row.to_dict()
    assert plain["reasoning_sections"]["review"] == "법리상 부당" and not plain["reasoning_sections"]["ai_label"]
    assert plain["legal_reasoning"] == "타당성 검토 결과 : 법리상 부당"
    row.review_text, row.ai_label, row.ai_opinions = "법리상 부당", "3개 모델 의견 일치", OPINIONS
    data = row.to_dict()
    assert [o["name"] for o in data["reasoning_sections"]["opinions"]] == ["Anthropic", "OpenAI", "Gemini"]
    assert data["legal_reasoning"].split("\n")[1] == "AI 교차검증 결과 참고 : 3개 모델 의견 일치"


def test_previous_results_without_structure_keep_their_text():
    old = {"legal_reasoning": "옛 형식 본문\n[AI 교차검토 참고 · 2개 모델 의견 일치]", "recommended_counteraction": "확인"}
    assert row_cell_text(old) == "옛 형식 본문\n[AI 교차검토 참고 · 2개 모델 의견 일치]\n\n대응 방안 : 확인"


def test_pdf_limit_cuts_each_part_not_the_model_opinions_away():
    sections = reasoning_sections("가" * 50, "3개 모델 의견 일치", OPINIONS)
    text = row_cell_text({"reasoning_sections": sections, "recommended_counteraction": "나" * 2000}, limit=900)
    assert "(Gemini : 부당)" in text and text.endswith(" …")


def test_docx_cell_keeps_line_breaks():
    from docx import Document

    doc = Document()
    cell = doc.add_table(rows=1, cols=1).rows[0].cells[0]
    cell.text = row_cell_text({"reasoning_sections": reasoning_sections("검토", "3개 모델 의견 일치", OPINIONS),
                               "recommended_counteraction": "원문 확인"})
    saved = io.BytesIO()
    doc.save(saved)
    reread = Document(io.BytesIO(saved.getvalue())).tables[0].rows[0].cells[0].text
    assert "타당성 검토 결과 : 검토\nAI 교차검증 결과 참고 : 3개 모델 의견 일치\n  (Anthropic : 부당)" in reread
    assert reread.endswith("\n\n대응 방안 : 원문 확인")


def test_web_table_renders_each_part_on_its_own_line():
    from playwright.sync_api import expect, sync_playwright

    static = ROOT / "apps/web/static"
    files = {f"/static/{p.relative_to(static).as_posix()}": p for p in static.rglob("*") if p.is_file()}
    files["/"] = ROOT / "apps/web/index.html"
    run = {"id": "r1", "project_id": "p1", "state": "COMPLETED", "progress": 1, "stage_message": "",
           "document_ids": ["d"], "scores": {}, "unverified_items": [], "input_snapshot": {"scope_revision": 0},
           "started_at": "2026-09-25T00:00:00"}
    project = {"id": "p1", "name": "합성 사건", "can_delete": True, "document_count": 1,
               "external_ai_policy": "MASKED", "scope_revision": 0}
    structured = {"location": "2면", "cited_authority": "대법원 2099다1 판결", "basis": "UNCONFIRMED", "claim_text": "c",
                  "ai_generation_basis": "b", "validity_verdict": "부당", "recommended_counteraction": "원문 확인",
                  "reasoning_sections": reasoning_sections("근거가 확장되지 않습니다.", "3개 모델 의견 불일치", OPINIONS),
                  "legal_reasoning": "무시되는 글"}
    legacy = {**structured, "location": "3면", "reasoning_sections": None, "legal_reasoning": "첫 줄\n둘째 줄"}
    result = {"documents": [{"filename": "a.pdf", "quarantined": False, "findings": [],
                             "engine_data": {"adversarial": {"scanned_layers": ["visible_text"], "adversarial_risk": "NONE"}},
                             "ai_hallucination_table": [structured, legacy]}]}

    def respond(route):
        path = urlsplit(route.request.url).path
        if path in files:
            return route.fulfill(path=str(files[path]))
        replies = {"/api/health": {"status": "ok", "version": "0.9.1"},
                   "/api/identity/me": {"user_id": "u", "role": "ADMIN", "authentication": "password"},
                   "/api/verification-runs": {"active_count": 0, "runs": []},
                   "/api/projects": [project], "/api/projects/p1": project, "/api/projects/p1/runs": [run]}
        if path in replies:
            return route.fulfill(json=replies[path])
        if path.endswith("/result"):
            return route.fulfill(json=result)
        if path.endswith("/case-matrix"):
            return route.fulfill(body="null", content_type="application/json")
        return route.fulfill(json=[])

    with sync_playwright() as playwright:
        options = {"headless": True}
        if os.getenv("LV_TEST_BROWSER_CHANNEL"):
            options["channel"] = os.environ["LV_TEST_BROWSER_CHANNEL"]
        browser = playwright.chromium.launch(**options)
        try:
            page = browser.new_page(viewport={"width": 1280, "height": 1400})
            page.route("**/*", respond)
            page.goto("http://layout.test/")
            page.wait_for_function("state.result && state.result.documents && state.result.documents.length === 1")
            page.evaluate("switchTab('ai-verification')")
            cell = page.locator("#aiVerificationRows tr").nth(0).locator("td").nth(3)
            lines = cell.locator(".reasoning-line")
            expect(lines).to_have_count(2)
            expect(lines.nth(0)).to_have_text("타당성 검토 결과 : 근거가 확장되지 않습니다.")
            expect(lines.nth(1)).to_have_text("AI 교차검증 결과 참고 : 3개 모델 의견 불일치")
            expect(cell.locator(".ai-opinion-verdict")).to_have_text(
                ["(Anthropic : 부당)", "(OpenAI : 타당)", "(Gemini : 부당)"])
            tops = cell.locator(".ai-opinion-list li").evaluate_all(
                "els => els.map(e => Math.round(e.getBoundingClientRect().top))")
            assert tops == sorted(tops) and len(set(tops)) == 3, "모델 의견은 줄마다 하나씩"
            assert cell.locator(".reasoning-line").nth(0).evaluate("e => e.getBoundingClientRect().bottom") <= \
                cell.locator(".reasoning-line").nth(1).evaluate("e => e.getBoundingClientRect().top")
            expect(cell).not_to_contain_text("무시되는 글")
            legacy_cell = page.locator("#aiVerificationRows tr").nth(1).locator("td").nth(3).locator(".legal-reasoning")
            assert legacy_cell.evaluate("e => getComputedStyle(e).whiteSpace") == "pre-wrap"
            page.close()
        finally:
            browser.close()
