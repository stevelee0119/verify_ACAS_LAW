"""Editable Word export from the same frozen payload as the generated report."""
from __future__ import annotations

import io
import json

from packages.common.terminology import EDITABLE_COPY_NOTICE, REVIEW_NOTICE
from .snapshot import json_lines, technical_payload, xml_text
from .summary import (FULL_RECORD_NOTE, SUMMARY, assessed_claims, detail_level, evidence_summary,
                      grouped_unverified, reviewed_workflow)


def build_report_docx(run_result, *, project=None, manifest=None, reveal_sealed=False):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    from docx.text.paragraph import Paragraph

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Header", "Footer"):
        style = doc.styles[name]
        style.font.name = "Malgun Gothic"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        style.font.size = Pt(11 if name == "Normal" else 14 if name.startswith("Heading") else 11)
    doc.styles["Title"].font.size = Pt(20)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.15
    metadata = getattr(run_result, "report_metadata", {})
    snapshot = getattr(run_result, "review_snapshot", {})
    doc.core_properties.title = "법률문서 검증보고서"
    doc.core_properties.author = str(metadata.get("created_by") or "ACASia_LAW")
    doc.core_properties.comments = "Editable copy; retained artifacts and snapshot hashes identify the generated version."

    body = doc.element.body
    # doc.add_paragraph는 매번 본문 전체에서 구역 설정(sectPr)을 찾아 그 앞에 넣는다.
    # 기술 부록처럼 문단이 수만 개면 비용이 제곱으로 늘어 Word 한 부에 4분이 걸렸다.
    # 구역 설정은 문서 끝의 같은 요소이므로 한 번만 찾아 두고 그 앞에 직접 넣는다.
    # 결과 XML은 같고 시간은 문단 수에 비례한다.
    section_properties = body.sectPr

    def paragraph(value, style=None):
        element = OxmlElement("w:p")
        if section_properties is not None:
            section_properties.addprevious(element)
        else:
            body.append(element)
        item = Paragraph(element, doc._body)
        text = xml_text(value)
        if text:
            item.add_run(text)
        if style is not None:
            item.style = style
        return item

    def table(headers, rows, widths):
        t = doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for column, width in zip(t.columns, widths):
            column.width = Inches(width)
        for cell, text in zip(t.rows[0].cells, headers):
            cell.text = xml_text(text)
        repeat = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            cells = t.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = xml_text(value if not isinstance(value, (dict, list)) else json.dumps(value, ensure_ascii=False, default=str))
        for index, row in enumerate(t.rows):
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                props = cell._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                for side in ("top", "left", "bottom", "right"):
                    edge = OxmlElement(f"w:{side}")
                    for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                        edge.set(qn(f"w:{key}"), value)
                    borders.append(edge)
                props.append(borders)
                margins = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    margin = OxmlElement(f"w:{side}")
                    margin.set(qn("w:w"), "90")
                    margin.set(qn("w:type"), "dxa")
                    margins.append(margin)
                props.append(margins)
                if index == 0:
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:fill"), "E7EDF0")
                    props.append(shading)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(3)
                    for run in p.runs:
                        run.font.size = Pt(10)
                        run.bold = index == 0
        paragraph("")

    doc.add_heading("법률문서 검증보고서", 0)
    paragraph(metadata.get("label", "DRAFT / 검토용 초안"), "Subtitle")
    project = project or {}
    paragraph(f"{project.get('name', '')} 사건의 문서 검증 결과와 사람의 검토 기록입니다. 시스템 판정은 실행 당시의 근거 범위에 한정됩니다.")
    paragraph(REVIEW_NOTICE)
    paragraph(EDITABLE_COPY_NOTICE)
    for key in ("source_run_id", "source_run_hash", "snapshot_hash", "export_snapshot_hash", "created_by", "created_at", "finalized_by", "finalized_at", "note"):
        if metadata.get(key) is not None:
            paragraph(f"{key}: {metadata[key]}")
    if not metadata:
        paragraph(f"source_run_id: {run_result.run_id}")
    summary = detail_level(run_result) == SUMMARY
    paragraph("보고서 분량: " + ("요약본 — 전체 기술 기록은 검증 상세(JSON)에 보존" if summary else "전체 기술 기록 포함"))
    doc.add_heading("검증 결과", 1)
    if summary:
        # 같은 제목·판정의 항목은 한 줄로 묶고 건수를 적는다. 설명은 중간 이상(참고 신호 제외)만 풀어 쓴다.
        groups = {}
        for f in run_result.all_findings:
            groups.setdefault((f.title, str(f.status), str(f.evidence_grade)), []).append(f)
        table(["항목", "시스템 판정", "근거 등급"],
              [[title + (f" ({len(items)}건)" if len(items) > 1 else ""), status, grade]
               for (title, status, grade), items in groups.items()], [4.5, 1.5, 1])
        detailed = [items[0] for items in groups.values()
                    if not items[0].advisory_only and str(items[0].severity) in ("CRITICAL", "HIGH", "MEDIUM")]
    else:
        table(["항목", "시스템 판정", "근거 등급"],
              [[f.title, str(f.status), str(f.evidence_grade)] for f in run_result.all_findings], [4.5, 1.5, 1])
        detailed = run_result.all_findings
    for finding in detailed:
        doc.add_heading(xml_text(finding.title), 2)
        paragraph(finding.detail)
    doc.add_heading("미검증 항목과 사용 불가 자료", 1)
    unverified = ([[i["target"] + (f" ({i['count']}건)" if i["count"] > 1 else ""), i["reason"]]
                   for i in grouped_unverified(run_result.unverified_items)] if summary else
                  [[item.get("raw_text") or item.get("document_id") or item.get("kind"), item.get("reason")]
                   for item in run_result.unverified_items])
    table(["대상", "사유"],
          unverified +
          [[s.get("name"), f"{s.get('status')} / {s.get('note', '')}"] for s in run_result.unavailable_sources], [2.5, 4.5])
    doc.add_heading("AI 작성 분석 및 법률 주장 타당성 검토", 1)
    ai_rows = []
    for d in run_result.documents:
        det = getattr(d, "ai_detector_result", {}) or {}
        a = d.authorship or {}
        v = det.get("verdict") or a.get("verdict", "-")
        score = f"{det.get('score'):.2f}" if "score" in det else str(a.get("score", "-"))
        reasons = "; ".join(det.get("reasons", [])) if det.get("reasons") else " ".join(a.get("notes", []))
        ai_rows.append([d.filename, v, score, reasons])
    if ai_rows:
        table(["문서", "AI 진단", "확신도", "판정 근거"], ai_rows, [1.5, 1.3, 0.8, 3.4])
    from .model_opinions import model_opinion_rows

    opinion_rows = model_opinion_rows(run_result.documents)
    if opinion_rows:
        doc.add_heading("모델별 AI 작성 판정", 2)
        table(["문서", "모델", "판정", "점수", "근거"], opinion_rows, [1.3, 1.2, 1.1, 0.5, 2.9])

    all_hallucination_rows = []
    for d in run_result.documents:
        for row in getattr(d, "ai_hallucination_table", []) or []:
            all_hallucination_rows.append([
                d.filename + " " + str(row.get("location", "")),
                str(row.get("claim_text", "")) + f"\n({row.get('cited_authority', '')})",
                str(row.get("ai_generation_basis", "")),
                str(row.get("legal_reasoning", "")) + f"\n[대응] {row.get('recommended_counteraction', '')}",
                str(row.get("validity_verdict", "")),
            ])
    if all_hallucination_rows:
        doc.add_heading("법률 인용 오류·근거 미확인 주장 대조표 (AI 작성 여부 판단과 별개)", 2)
        table(["위치", "문서 주장 / 인용", "AI 생성 근거", "법리적 검토 및 반박 근거", "평가"], all_hallucination_rows, [1.0, 1.8, 1.5, 2.0, 0.7])

    doc.add_heading("사람의 검토 기록", 1)
    workflow = snapshot.get("workflow", [])
    shown, untouched = reviewed_workflow(workflow) if summary else (workflow, 0)
    table(["항목", "진행과 의견", "검토자와 메모"],
          [[r.get("title") or r.get("finding_id"), f"{r.get('workflow_state')} / {r.get('decision')}",
            f"{r.get('updated_by', '')} / {r.get('updated_at', '')}\n{r.get('note', '')}"]
           for r in shown], [2.5, 1.8, 2.7])
    if untouched:
        paragraph(f"검토를 시작하지 않은 항목 {untouched}건은 목록에서 생략했습니다(전체 목록은 검증 상세 JSON).")
    doc.add_heading("쟁점과 주장 및 증거 관계", 1)
    for issue in snapshot.get("matrix", {}).get("issues", []):
        doc.add_heading(xml_text(issue.get("title")), 2)
        paragraph(f"요건사실: {issue.get('elements', [])}")
        paragraph(f"법적 근거: {issue.get('legal_basis', '')} / 기준일: {issue.get('reference_date')}")
    rows = []
    claims = snapshot.get("matrix", {}).get("claims", [])
    shown, unassessed = assessed_claims(claims) if summary else (claims, 0)
    for item in shown:
        assessment = item.get("assessment", {})
        rows.append([item["claim"].get("text"), f"{assessment.get('position', 'UNASSESSED')} / {item.get('review_status')}",
                     {"issue_id": assessment.get("issue_id"), "evidence_links": assessment.get("evidence_links", []),
                      "missing_material": assessment.get("missing_material", "")}])
    table(["주장", "입장과 증거 검토", "쟁점 및 증거 연결"], rows, [2.5, 1.8, 2.7])
    if unassessed:
        paragraph(f"입장·증거를 아직 적지 않은 주장 {unassessed}건은 목록에서 생략했습니다(전체 목록은 검증 상세 JSON).")
    if summary:
        evidence = evidence_summary(run_result)
        doc.add_heading("검증 근거 요약", 1)
        table(["위치", "수행하지 못한 검사 단계"],
              [[s["path"], str(s["stage"])] for s in evidence["unavailable_stages"]] or [["-", "기록된 항목 없음"]],
              [2.5, 4.5])
        table(["출처", "상태", "건수"], [[c["adapter"], c["status"], str(c["count"])] for c in evidence["source_counts"]]
              or [["-", "-", "0"]], [3, 2.5, 1.5])
        if evidence["problem_sources"]:
            table(["출처 / 상태", "조회 / 주소", "비고"],
                  [[f"{p['adapter']} / {p['status']}", f"{p.get('query') or ''}\n{p.get('url') or ''}",
                    f"{p.get('note') or ''}\n{p.get('retrieved_at') or ''}"] for p in evidence["problem_sources"]],
                  [1.8, 3.2, 2.0])
            if evidence["problem_sources_omitted"]:
                paragraph(f"성공하지 못한 조회 외 {evidence['problem_sources_omitted']}건은 검증 상세 JSON에 있습니다.")
        doc.add_heading("전체 기술 기록", 1)
        paragraph(FULL_RECORD_NOTE)
        footer = section.footer.paragraphs[0]
        footer.text = xml_text(metadata.get("label", "DRAFT") + " | " + run_result.run_id)
        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()
    doc.add_heading("전체 기술 부록", 1)
    paragraph("실행 당시 기록된 검사 상태, 수행하지 못한 단계, 전체 출처와 모델 실행 기록을 포함합니다. 기록이 없다는 사실은 검사 성공을 뜻하지 않습니다. 공유용에서 제외된 자료는 공유 정책에 표시됩니다.")
    payload = technical_payload(run_result)
    if manifest:
        payload["export_manifest"] = manifest
    for path, value in json_lines(payload):
        p = paragraph(path)
        p.paragraph_format.keep_with_next = True
        for run in p.runs:
            run.bold = True
            run.font.size = Pt(9)
        for offset in range(0, max(1, len(value)), 3000):
            paragraph(value[offset:offset + 3000])
    footer = section.footer.paragraphs[0]
    footer.text = xml_text(metadata.get("label", "DRAFT") + " | " + run_result.run_id)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
