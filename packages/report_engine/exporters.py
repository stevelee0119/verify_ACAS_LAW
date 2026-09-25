"""제20.2장 출력 파일: JSON / CSV / XLSX / Chain of Custody Manifest."""
from __future__ import annotations

import csv
import io
import json
from datetime import datetime
from typing import Any, Dict, List

from packages.common.enums import MM4_ADVISORY_TYPES, spec_code
from packages.common.terminology import TERMINOLOGY_VERSION, terminology_catalog, EDITABLE_COPY_NOTICE
from .serialize import dumps, jsonable_payload
from .snapshot import compact_claim_rows, json_lines, xml_text

FINDING_COLUMNS = [
    "finding_id", "type", "spec_code", "status", "severity", "evidence_grade", "confidence",
    "document_id", "page", "block_id", "title", "detail", "engine",
    "meta_message_type", "advisory_only", "review_status", "sources", "has_sealed_content",
]


def findings_to_rows(findings: List[Any], *, reveal_sealed: bool = False) -> List[Dict[str, Any]]:
    rows = []
    for finding in findings:
        data = finding.to_dict(reveal_sealed=reveal_sealed)
        row = {key: data.get(key) for key in FINDING_COLUMNS}
        # 명세 제1.2장의 코드명. 이 시스템의 이름과 다른 경우에만 값이 달라진다.
        # 같은 판정 유형 안의 세부 결함(예: COURT_CODE_MISMATCH, CROSS_DOC_INCONSISTENCY)은 그 이름을 우선한다.
        row["spec_code"] = (getattr(finding, "confidence_features", None) or {}).get("defect_code") or spec_code(finding.type)
        row["sources"] = ";".join(data.get("sources") or [])
        rows.append(row)
    return rows


def to_json(run_result: Any, *, reveal_sealed: bool = False) -> bytes:
    """제20.2장 JSON 전체 검증결과(바이트)."""
    return dumps(to_payload(run_result, reveal_sealed=reveal_sealed), indent=2).encode("utf-8")


def to_payload(run_result: Any, *, reveal_sealed: bool = False) -> dict:
    """같은 내용을 파이썬 객체로 돌려준다.

    DB의 JSON 칼럼에 넣을 때는 문자열로 만들었다가 다시 읽을 이유가 없다.
    큰 결과에서 그 왕복은 수 초가 걸리고, json.dumps/loads는 그동안 GIL을
    놓지 않아 임차 갱신 스레드가 밀린다.
    """
    payload = {
        "product": "ACASia_LAW",
        "terminology_version": TERMINOLOGY_VERSION,
        "terminology": terminology_catalog(),
        "run_id": run_result.run_id,
        "project_id": run_result.project_id,
        "state": str(run_result.state),
        "verification_key": run_result.verification_key,
        "started_at": run_result.started_at.isoformat(),
        "finished_at": run_result.finished_at.isoformat() if run_result.finished_at else None,
        "scores": run_result.scores,
        "timeline": run_result.timeline,
        "unavailable_sources": run_result.unavailable_sources,
        "unverified_items": run_result.unverified_items,
        "errors": run_result.errors,
        "input_snapshot": getattr(run_result, "input_snapshot", {}),
        # 엔진별 실행 여부·입력·finding 수·시간·건너뛴 사유(v4 P0). '0건'이 '실행했는데 없음'인지 가른다.
        "run_manifest": getattr(run_result, "run_manifest", {}) or {},
        "documents": [
            {
                "document_id": d.document_id,
                "filename": d.filename,
                "sha256": d.normalized.sha256 if d.normalized else None,
                "parser": d.normalized.parser_name if d.normalized else None,
                "quarantined": d.quarantined,
                "rag_indexable": d.rag_indexable,
                "warnings": d.warnings,
                "authorship": d.authorship,
                "ai_detector_result": getattr(d, "ai_detector_result", {}),
                "ai_hallucination_table": getattr(d, "ai_hallucination_table", []),
                "argument_validity_summary": getattr(d, "argument_validity_summary", ""),
                "masked_preview": d.masked_preview,
                "citations": d.citations,
                "claims": d.claims,
                "entities": d.entities,
                "events": d.events,
                "engine_data": d.engine_data,
                "source_records": [r.to_dict() if hasattr(r, "to_dict") else r for r in d.source_records],
                "pages": ([{"page_number": p.page_number, "width": p.width, "height": p.height,
                            "blocks": [b.to_dict() for b in p.blocks if b.visible and b.source_layer in ("visible_text", "ocr_layer")]}
                           for p in d.normalized.pages] if d.normalized and hasattr(d.normalized, "pages")
                          else getattr(d, "pages", [])),
                "page_coverage": d.engine_data.get("page_coverage", []),
                "findings": [f.to_dict(reveal_sealed=reveal_sealed) for f in d.findings],
            }
            for d in run_result.documents
        ],
        "project_findings": [f.to_dict(reveal_sealed=reveal_sealed) for f in run_result.project_findings],
        "model_executions": getattr(run_result, "model_executions", []),
        "notice": (
            "봉인된 원문(sealed)은 사용자 명시적 열람 요청이 있는 경우에만 포함된다. "
            "본 결과는 검증 보조자료이며 최종 법적 평가는 사용자에게 있다."
        ),
    }
    if hasattr(run_result, "report_metadata"):
        payload["report"] = run_result.report_metadata
        payload["review_snapshot"] = {key: value for key, value in run_result.review_snapshot.items()
                                      if key != "engine_result"}
    return payload


def safe_cell(value):
    if isinstance(value, (dict, list, tuple)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, str):
        value = xml_text(value)
        if value.lstrip(" \t\r\n\ufeff").startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n")):
            return "'" + value
    return value


def to_csv(findings: List[Any], *, reveal_sealed: bool = False, report_metadata=None) -> bytes:
    """제20.2장 CSV Findings."""
    buffer = io.StringIO()
    metadata = report_metadata or {}
    extra = {"report_state": metadata.get("state"), "report_label": metadata.get("label"),
             "report_audience": metadata.get("audience"), "source_run_id": metadata.get("source_run_id"),
             "snapshot_hash": metadata.get("snapshot_hash"), "created_by": metadata.get("created_by"),
             "finalized_by": metadata.get("finalized_by"), "finalized_at": metadata.get("finalized_at"),
             "editable_copy_notice": EDITABLE_COPY_NOTICE} if report_metadata else {}
    writer = csv.DictWriter(buffer, fieldnames=FINDING_COLUMNS + list(extra), extrasaction="ignore")
    writer.writeheader()
    rows = findings_to_rows(findings, reveal_sealed=reveal_sealed)
    for row in rows or ([{}] if extra else []):
        writer.writerow({key: safe_cell(value) for key, value in (row | extra).items()})
    return buffer.getvalue().encode("utf-8-sig")  # Excel 한글 호환


def to_xlsx(run_result: Any, *, reveal_sealed: bool = False) -> bytes:
    """제20.2장 XLSX 오류표."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = openpyxl.Workbook()
    overflow = None

    def append(worksheet, values):
        nonlocal overflow
        row = []
        for index, raw in enumerate(values, start=1):
            value = safe_cell(raw)
            if isinstance(value, str) and len(value) > 30000:
                if overflow is None:
                    overflow = workbook.create_sheet("Long values")
                    overflow.append(["sheet", "row", "column", "part", "text"])
                for part, offset in enumerate(range(0, len(value), 30000), start=1):
                    overflow.append([worksheet.title, worksheet.max_row + 1, index, part, safe_cell(value[offset:offset + 30000])])
                value = f"[Full value in Long values: {worksheet.title}, row {worksheet.max_row + 1}, column {index}]"
            row.append(value)
        worksheet.append(row)

    # 1) 요약
    summary = workbook.active
    summary.title = "검증개요"
    summary.append(["항목", "값"])
    summary.append(["Run ID", run_result.run_id])
    summary.append(["Project ID", run_result.project_id])
    summary.append(["상태", str(run_result.state)])
    summary.append(["Verification Key", run_result.verification_key])
    summary.append(["생성시각", datetime.utcnow().isoformat()])
    summary.append(["문서 수", len(run_result.documents)])
    summary.append(["Finding 수", len(run_result.all_findings)])
    # 제9.2장 배포가능 상태. 보고서를 여는 사람이 가장 먼저 봐야 하는 값이다.
    gate = (run_result.scores or {}).get("release_gate") or {}
    if gate:
        append(summary, ["배포가능 상태", gate.get("release_gate")])
        append(summary, ["검증위험 지수", gate.get("hallucination_risk")])
        append(summary, ["지수 성격", gate.get("risk_index_note")])
        if gate.get("hard_block_reasons"):
            append(summary, ["차단 사유", "\n".join(gate["hard_block_reasons"])])
        if gate.get("review_reasons"):
            append(summary, ["사람 검토 사유", "\n".join(gate["review_reasons"])])
    for key, value in getattr(run_result, "report_metadata", {}).items():
        append(summary, [key, value])
    append(summary, ["editable_copy_notice", EDITABLE_COPY_NOTICE])
    for axis, value in (run_result.scores.get("axes") or {}).items():
        summary.append([axis, json.dumps(value, ensure_ascii=False)])
    summary.append(["사용 불가 Source", ", ".join(s["name"] for s in run_result.unavailable_sources)])

    # 2) Findings
    sheet = workbook.create_sheet("Findings")
    sheet.append(FINDING_COLUMNS)
    severity_fill = {
        "CRITICAL": PatternFill("solid", fgColor="FFC7CE"),
        "HIGH": PatternFill("solid", fgColor="FFD9B3"),
        "MEDIUM": PatternFill("solid", fgColor="FFF2CC"),
    }
    main = [f for f in run_result.all_findings if not f.advisory_only and f.type not in MM4_ADVISORY_TYPES]
    for finding in sorted(main, key=lambda f: -f.severity.rank):
        row = findings_to_rows([finding], reveal_sealed=reveal_sealed)[0]
        append(sheet, [row.get(c) for c in FINDING_COLUMNS])
        fill = severity_fill.get(str(finding.severity))
        if fill:
            for cell in sheet[sheet.max_row]:
                cell.fill = fill

    # 3) 참고 신호 (MM-4) — 본문과 분리 (제7-A.4장)
    advisory_sheet = workbook.create_sheet("참고신호(MM-4)")
    advisory_sheet.append(FINDING_COLUMNS)
    for finding in run_result.all_findings:
        if finding.advisory_only or finding.type in MM4_ADVISORY_TYPES:
            row = findings_to_rows([finding], reveal_sealed=reveal_sealed)[0]
            append(advisory_sheet, [row.get(c) for c in FINDING_COLUMNS])

    # 4) 미검증 항목
    unverified = workbook.create_sheet("미검증항목")
    unverified.append(["kind", "document_id", "citation_id", "type", "raw_text", "reason"])
    for item in run_result.unverified_items:
        append(unverified, [item.get(k) for k in ("kind", "document_id", "citation_id", "type", "raw_text", "reason")])

    snapshot = getattr(run_result, "review_snapshot", {})
    review_sheet = workbook.create_sheet("검토기록")
    columns = ["finding_id", "system_status", "workflow_state", "decision", "assignee", "note", "updated_by", "updated_at", "revision"]
    append(review_sheet, columns)
    for row in snapshot.get("workflow", []):
        append(review_sheet, [row.get(c) for c in columns])
    matrix = workbook.create_sheet("쟁점증거표")
    append(matrix, ["claim_id", "document_id", "claim", "issue_id", "position", "support_status", "evidence_links", "missing_material", "note"])
    for row in snapshot.get("matrix", {}).get("claims", []):
        assessment = row.get("assessment", {})
        append(matrix, [row["claim"].get("claim_id"), row.get("document_id"), row["claim"].get("text"),
                        assessment.get("issue_id"), assessment.get("position"), row.get("review_status"),
                        assessment.get("evidence_links", []), assessment.get("missing_material"), assessment.get("note")])
    issues = workbook.create_sheet("쟁점")
    columns = ["id", "title", "elements", "reference_date", "legal_basis", "priority", "note"]
    append(issues, columns)
    for row in snapshot.get("matrix", {}).get("issues", []):
        append(issues, [row.get(c) for c in columns])
    from .summary import FULL_RECORD_NOTE, SUMMARY, detail_level, evidence_summary

    if detail_level(run_result) == SUMMARY:
        # 요약본: 수행하지 못한 단계와 실패한 조회는 빠짐없이, 성공한 조회는 건수로 싣는다.
        evidence = evidence_summary(run_result)
        basis = workbook.create_sheet("검증근거")
        append(basis, ["구분", "대상", "상태", "내용"])
        for stage in evidence["unavailable_stages"]:
            append(basis, ["수행하지 못한 단계", stage["path"], "UNAVAILABLE", stage["stage"]])
        for count in evidence["source_counts"]:
            append(basis, ["출처 조회 건수", count["adapter"], count["status"], count["count"]])
        for problem in evidence["problem_sources"]:
            append(basis, ["성공하지 못한 조회", problem["adapter"], problem["status"],
                           f"{problem.get('query') or ''} {problem.get('url') or ''} {problem.get('note') or ''}".strip()])
        append(basis, ["전체 기술 기록", "검증 상세(JSON)", "", FULL_RECORD_NOTE])
    else:
        technical = workbook.create_sheet("기술부록")
        append(technical, ["JSON path", "part", "value"])
        appendix = jsonable_payload(to_payload(run_result, reveal_sealed=reveal_sealed))
        for path, value in json_lines(compact_claim_rows(appendix)):
            for part, offset in enumerate(range(0, max(1, len(value)), 29000), start=1):
                append(technical, [path, part, value[offset:offset + 29000]])

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        for row in worksheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = safe_cell(cell.value)
                    cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="center")
        for column in worksheet.columns:
            width = max((len(str(c.value)) for c in column if c.value), default=10)
            worksheet.column_dimensions[column[0].column_letter].width = min(60, max(12, width + 2))

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def to_manifest(audit_chain: Any, project_id: str, run_result: Any) -> bytes:
    """제20.2장 Chain of Custody Manifest JSON."""
    manifest = audit_chain.manifest(project_id=project_id)
    manifest["run"] = {
        "run_id": run_result.run_id,
        "verification_key": run_result.verification_key,
        "state": str(run_result.state),
        "documents": [
            {
                "document_id": d.document_id,
                "filename": d.filename,
                "sha256": d.normalized.sha256 if d.normalized else None,
            }
            for d in run_result.documents
        ],
    }
    return json.dumps(manifest, ensure_ascii=False, indent=2, default=str).encode("utf-8")
