"""TXT / CSV / XLS·XLSX / 이미지 Parser."""
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List

from packages.common.schemas import BBox, Block, NormalizedDocument, Page, new_id

from .base import DocumentParser, ParserError
from .ocr import get_ocr_adapter


class TextParser(DocumentParser):
    name = "TextParser"
    extensions = [".txt", ".md", ".log", ".json", ".xml", ".html", ".htm"]
    mime_types = ["text/plain", "text/markdown", "application/json", "text/html"]

    def parse(self, path: str, *, document_id: str, filename: str, mime_type: str, sha256: str) -> NormalizedDocument:
        raw = Path(path).read_bytes()
        text = raw.decode("utf-8", "replace")
        doc = NormalizedDocument(
            document_id=document_id,
            filename=filename,
            mime_type=mime_type or "text/plain",
            sha256=sha256,
            parser_name=self.name,
        )
        page = Page(page_number=1)
        for line in text.splitlines():
            if line.strip():
                page.blocks.append(Block(block_id=new_id("B"), text=line.strip(), page=1))
        doc.pages.append(page)
        doc.raw_layers["rendered_text"] = text
        doc.raw_layers["raw_text"] = text
        return doc


class SpreadsheetParser(DocumentParser):
    name = "SpreadsheetParser"
    extensions = [".xlsx", ".xlsm", ".xltx", ".csv", ".tsv", ".xls"]
    mime_types = [
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "text/csv",
        "application/vnd.ms-excel",
    ]

    def parse(self, path: str, *, document_id: str, filename: str, mime_type: str, sha256: str) -> NormalizedDocument:
        doc = NormalizedDocument(
            document_id=document_id,
            filename=filename,
            mime_type=mime_type or "text/csv",
            sha256=sha256,
            parser_name=self.name,
        )
        suffix = Path(filename).suffix.lower()
        page = Page(page_number=1)
        lines: List[str] = []
        hidden_lines: List[str] = []

        if suffix in (".csv", ".tsv"):
            import csv

            delimiter = "\t" if suffix == ".tsv" else ","
            with open(path, newline="", encoding="utf-8", errors="replace") as f:
                for row_index, row in enumerate(csv.reader(f, delimiter=delimiter), start=1):
                    text = " | ".join(row)
                    if text.strip():
                        page.blocks.append(
                            Block(block_id=new_id("B"), text=text, page=1, block_type="table",
                                  attributes={"row": row_index})
                        )
                        lines.append(text)
        else:
            try:
                import openpyxl
            except ImportError as exc:  # pragma: no cover
                raise ParserError(f"openpyxl 미설치: {exc}") from exc
            try:
                wb = openpyxl.load_workbook(path, data_only=False)
            except Exception as exc:
                raise ParserError(f"스프레드시트 열기 실패: {exc}") from exc

            hidden_sheets: List[str] = []
            hidden_rows: List[Dict[str, Any]] = []
            formulas: List[Dict[str, Any]] = []
            defined_names: List[str] = []
            try:
                defined_names = list(wb.defined_names.keys())  # type: ignore[attr-defined]
            except Exception:
                pass

            for ws in wb.worksheets:
                sheet_hidden = ws.sheet_state != "visible"
                if sheet_hidden:
                    hidden_sheets.append(ws.title)
                for dim_key, dim in list(getattr(ws, "row_dimensions", {}).items()):
                    if getattr(dim, "hidden", False):
                        hidden_rows.append({"sheet": ws.title, "row": dim_key})
                for col_key, dim in list(getattr(ws, "column_dimensions", {}).items()):
                    if getattr(dim, "hidden", False):
                        hidden_rows.append({"sheet": ws.title, "column": col_key})
                for row in ws.iter_rows():
                    values = []
                    row_hidden = getattr(ws.row_dimensions.get(row[0].row), "hidden", False) if row else False
                    for cell in row:
                        v = cell.value
                        if v is None:
                            continue
                        s = str(v)
                        values.append(s)
                        if s.startswith("="):
                            formulas.append({"sheet": ws.title, "cell": cell.coordinate, "formula": s})
                        font = getattr(cell, "font", None)
                        color = getattr(getattr(font, "color", None), "rgb", None) if font else None
                        if color and str(color).upper() in ("FFFFFFFF", "00FFFFFF"):
                            hidden_lines.append(s)
                            page.blocks.append(
                                Block(
                                    block_id=new_id("B"),
                                    text=s,
                                    page=1,
                                    source_layer="hidden_text",
                                    visible=False,
                                    attributes={"sheet": ws.title, "cell": cell.coordinate, "hidden_reason": "WHITE_FONT"},
                                )
                            )
                    if not values:
                        continue
                    text = " | ".join(values)
                    is_hidden = sheet_hidden or bool(row_hidden)
                    page.blocks.append(
                        Block(
                            block_id=new_id("B"),
                            text=text,
                            page=1,
                            block_type="table",
                            source_layer="hidden_text" if is_hidden else "visible_text",
                            visible=not is_hidden,
                            attributes={"sheet": ws.title, "row": row[0].row,
                                        "hidden_reason": "HIDDEN_SHEET_OR_ROW" if is_hidden else None},
                        )
                    )
                    (hidden_lines if is_hidden else lines).append(text)

            doc.structure.update(
                {
                    "hidden_sheets": hidden_sheets,
                    "hidden_rows": hidden_rows,
                    "formulas": formulas,
                    "defined_names": defined_names,
                    "sheet_names": [ws.title for ws in wb.worksheets],
                }
            )
            props = wb.properties
            for key in ("creator", "lastModifiedBy", "title", "subject", "keywords", "description", "company"):
                value = getattr(props, key, None)
                if value:
                    doc.metadata[key] = str(value)

        doc.pages.append(page)
        doc.raw_layers["rendered_text"] = "\n".join(lines)
        doc.raw_layers["raw_text"] = "\n".join(lines + hidden_lines)
        doc.raw_layers["hidden_text"] = "\n".join(hidden_lines)
        doc.raw_layers["metadata_text"] = "\n".join(f"{k}: {v}" for k, v in doc.metadata.items())
        return doc


EXIF_GPS_TAG = 34853


class ImageParser(DocumentParser):
    name = "ImageParser"
    extensions = [".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tif", ".tiff", ".webp"]
    mime_types = ["image/jpeg", "image/png", "image/gif", "image/bmp", "image/tiff", "image/webp"]

    def parse(self, path: str, *, document_id: str, filename: str, mime_type: str, sha256: str) -> NormalizedDocument:
        doc = NormalizedDocument(
            document_id=document_id,
            filename=filename,
            mime_type=mime_type or "image/jpeg",
            sha256=sha256,
            parser_name=self.name,
        )
        try:
            from PIL import Image, ExifTags
        except ImportError as exc:  # pragma: no cover
            raise ParserError(f"Pillow 미설치: {exc}") from exc

        with Image.open(path) as im:
            width, height = im.size
            doc.structure.update(
                {
                    "image_size": [width, height],
                    "image_mode": im.mode,
                    "image_format": im.format,
                    "dpi": list(im.info.get("dpi", ())) or None,
                    "has_icc_profile": bool(im.info.get("icc_profile")),
                }
            )
            exif_data: Dict[str, Any] = {}
            try:
                exif = im.getexif()
                for tag_id, value in exif.items():
                    tag = ExifTags.TAGS.get(tag_id, str(tag_id))
                    exif_data[str(tag)] = str(value)[:500]
                gps = exif.get_ifd(EXIF_GPS_TAG) if hasattr(exif, "get_ifd") else None
                if gps:
                    doc.structure["gps"] = {ExifTags.GPSTAGS.get(k, str(k)): str(v)[:100] for k, v in gps.items()}
            except Exception as exc:  # pragma: no cover
                doc.parse_warnings.append(f"EXIF 추출 실패: {exc}")
            doc.structure["exif"] = exif_data
            doc.metadata.update({f"exif:{k}": v for k, v in exif_data.items()})
            if "thumbnail" in im.info or exif_data.get("JPEGInterchangeFormat"):
                doc.structure["has_embedded_thumbnail"] = True
            try:
                if im.format == "PNG" and getattr(im, "text", None):
                    doc.structure["png_text_chunks"] = {k: str(v)[:2000] for k, v in im.text.items()}
                    doc.metadata.update({f"png:{k}": str(v)[:500] for k, v in im.text.items()})
            except Exception:
                pass

        page = Page(page_number=1, width=float(width), height=float(height))
        ocr = get_ocr_adapter()
        ocr_lines = ocr.recognize(path)
        for line in ocr_lines:
            page.blocks.append(
                Block(
                    block_id=new_id("B"),
                    text=line.text,
                    page=1,
                    bbox=line.bbox,
                    source_layer="ocr_layer",
                    attributes={"ocr_confidence": line.confidence, "ocr_engine": ocr.name},
                )
            )
        doc.pages.append(page)
        ocr_text = "\n".join(line.text for line in ocr_lines)
        doc.raw_layers["ocr_layer"] = ocr_text
        doc.raw_layers["rendered_text"] = ocr_text
        doc.raw_layers["raw_text"] = ocr_text
        doc.raw_layers["metadata_text"] = "\n".join(f"{k}: {v}" for k, v in doc.metadata.items())
        if not ocr.available:
            doc.parse_warnings.append(
                "OCR Adapter가 사용 불가하여 이미지 본문을 추출하지 못했다. 관련 검증 항목은 UNVERIFIED로 표시한다."
            )
        return doc
